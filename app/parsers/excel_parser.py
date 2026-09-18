from .base import BaseParser, RawTable


class ExcelParser(BaseParser):
    """openpyxl read_only streaming for xlsx/xlsm; xlrd for legacy .xls."""
    engine = "openpyxl-stream"

    def _sheets_wanted(self, names):
        ids = [s for s in (self.req.sheetIdentifiers or []) if str(s).strip()]
        if not ids:
            return names[:1]                          # default: first sheet
        picked = []
        for i in ids:
            if str(i).isdigit() and int(i) < len(names):
                picked.append(names[int(i)])
            elif i in names:
                picked.append(i)
            else:
                self.warnings.append(f"sheet '{i}' not found; available: {names}")
        return picked

    def tables(self, path):
        header_row = max(int(self.req.headersRow or 1), 1)
        if path.lower().endswith(".xls"):
            yield from self._xls(path, header_row)
            return
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)   # data_only -> cached formula values
        try:
            for name in self._sheets_wanted(wb.sheetnames):
                it = wb[name].iter_rows(values_only=True)
                for _ in range(header_row - 1):
                    next(it, None)
                headers = next(it, None)
                if headers is None:
                    continue
                # trim trailing empty header cells (openpyxl reports max used column)
                last = max((i for i, h in enumerate(headers) if h not in (None, "")), default=-1)
                headers = list(headers[: last + 1])
                n = len(headers)

                def rows(it=it, n=n):
                    for r in it:
                        r = list(r[:n])
                        if any(v not in (None, "") for v in r):
                            yield r
                yield RawTable(headers, rows(), header_row + 1, {"sheet": name})
        finally:
            wb.close()     # runs after the consumer has drained the last sheet's rows

    def _xls(self, path, header_row):
        import xlrd
        book = xlrd.open_workbook(path)
        for name in self._sheets_wanted(book.sheet_names()):
            sh = book.sheet_by_name(name)
            if sh.nrows < header_row:
                continue
            headers = sh.row_values(header_row - 1)
            def rows(sh=sh):
                for i in range(header_row, sh.nrows):
                    r = sh.row_values(i)
                    if any(v not in (None, "") for v in r):
                        yield r
            yield RawTable(headers, rows(), header_row + 1, {"sheet": name})
